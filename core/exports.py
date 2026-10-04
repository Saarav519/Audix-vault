"""Letterhead helpers for Excel (openpyxl) and PDF (ReportLab) exports."""

from __future__ import annotations

import io
import os
from datetime import datetime
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.formatting import MINUS, RUPEE, fmt_date

FIRM = "Audix Solutions & Co"
TAGLINE = "Where Audits Meet Exceptionalism"
ASSOCIATED = "Associated with Vikas Kshitij & Associates, Chartered Accountants"

# Indian digit grouping in Excel (12,34,567). Negative values get a leading minus.
INR_FORMAT = '[>=10000000]"₹"##\\,##\\,##\\,##0;[>=100000]"₹"##\\,##\\,##0;"₹"#,##0'
INR_FORMAT_NEG = '"₹"#,##0;-"₹"#,##0'
QTY_FORMAT = "#,##0.###"
PCT_FORMAT = "0.00%"


def inr_number_format(value) -> str:
    try:
        return INR_FORMAT if Decimal(value) >= 0 else INR_FORMAT_NEG
    except Exception:
        return INR_FORMAT


# ---------------------------------------------------------------- Excel


class Sheet:
    def __init__(self, ws, client_name: str, title: str, period: str):
        self.ws = ws
        lime = PatternFill("solid", fgColor="B9E10C")
        dark = PatternFill("solid", fgColor="12170F")
        ws.sheet_view.showGridLines = False
        rows = [
            ("Audix Vault", Font(bold=True, size=16, color="FFFFFF"), dark),
            (FIRM, Font(bold=True, size=11, color="FFFFFF"), dark),
            (TAGLINE, Font(italic=True, size=10, color="B9E10C"), dark),
            (f"{title} · {client_name}", Font(bold=True, size=12), lime),
            (f"Period: {period}", Font(size=10), None),
            (f"Generated: {fmt_date(timezone.localdate())} {timezone.localtime():%H:%M} IST", Font(size=10, color="5A6556"), None),
        ]
        for i, (text, font, fill) in enumerate(rows, start=1):
            c = ws.cell(row=i, column=1, value=text)
            c.font = font
            if fill:
                for col in range(1, 12):
                    ws.cell(row=i, column=col).fill = fill
        self.row = len(rows) + 2

    def header(self, labels, widths=None):
        thin = Side(style="thin", color="DADFD2")
        for i, label in enumerate(labels, start=1):
            c = self.ws.cell(row=self.row, column=i, value=label)
            c.font = Font(bold=True, color="12180F")
            c.fill = PatternFill("solid", fgColor="F7F9F3")
            c.border = Border(bottom=thin)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if widths:
                self.ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        self.ws.freeze_panes = self.ws.cell(row=self.row + 1, column=1)
        self.row += 1

    def add(self, values, kinds=None, bold=False):
        """kinds per column: 'inr', 'qty', 'pct' (value is a percentage number, stored as a fraction), 'date', None."""
        for i, v in enumerate(values, start=1):
            kind = kinds[i - 1] if kinds else None
            if kind == "pct" and v is not None:
                v = float(Decimal(v) / 100)
            elif kind in ("inr", "qty") and v is not None:
                v = float(v)
            c = self.ws.cell(row=self.row, column=i, value=v)
            if kind == "inr" and v is not None:
                c.number_format = inr_number_format(v)
            elif kind == "qty":
                c.number_format = QTY_FORMAT
            elif kind == "pct":
                c.number_format = PCT_FORMAT
            elif kind == "date":
                c.number_format = "d mmm yyyy"
            elif isinstance(v, str):
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if bold:
                c.font = Font(bold=True)
        self.row += 1

    def footer(self):
        self.row += 1
        c = self.ws.cell(row=self.row, column=1, value=ASSOCIATED)
        c.font = Font(italic=True, size=9, color="5A6556")
        self.ws.cell(row=self.row + 1, column=1, value=f"Powered by {FIRM}").font = Font(size=9, color="5A6556")


def new_workbook():
    wb = Workbook()
    return wb, wb.active


def workbook_bytes(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- PDF

_FONT = None
BUNDLED_FONT = os.path.join(os.path.dirname(__file__), "fonts", "DejaVuSans.ttf")
# DejaVu Sans is bundled (core/fonts) so ₹ and − print on servers without system fonts, such as Railway.
FONT_PATHS = [
    os.environ.get("PDF_FONT_PATH", ""),
    BUNDLED_FONT,
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
]


def pdf_fonts() -> tuple[str, str, bool]:
    """(regular, bold, has_rupee_glyph). Falls back to Helvetica with 'Rs' when no TTF is available."""
    global _FONT
    if _FONT is None:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        _FONT = ("Helvetica", "Helvetica-Bold", False)
        for p in FONT_PATHS:
            if p and os.path.exists(p):
                try:
                    pdfmetrics.registerFont(TTFont("AudixSans", p))
                    bold = p.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
                    if os.path.exists(bold):
                        pdfmetrics.registerFont(TTFont("AudixSans-Bold", bold))
                        _FONT = ("AudixSans", "AudixSans-Bold", True)
                    else:
                        _FONT = ("AudixSans", "AudixSans", True)
                    break
                except Exception:
                    continue
    return _FONT


def pdf_text(text) -> str:
    """Make text safe for the active PDF font."""
    text = "" if text is None else str(text)
    if not pdf_fonts()[2]:
        text = text.replace(RUPEE, "Rs ").replace(MINUS, "-").replace("▼", "v").replace("▲", "^")
    return text


def build_pdf(story_fn, client_name: str, title: str, period: str, landscape_mode=True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate

    regular, bold, _ = pdf_fonts()
    pagesize = landscape(A4) if landscape_mode else A4
    buf = io.BytesIO()
    generated = f"Generated {fmt_date(timezone.localdate())} {timezone.localtime():%H:%M} IST"

    def letterhead(canvas, doc):
        w, h = pagesize
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#12170F"))
        canvas.rect(0, h - 26 * mm, w, 26 * mm, stroke=0, fill=1)
        # wordmark: "Audi" + two-tone x + "Vault"
        canvas.setFillColor(colors.white)
        canvas.setFont(bold, 18)
        x0, y0 = 14 * mm, h - 14 * mm
        canvas.drawString(x0, y0, "Audi")
        xw = x0 + canvas.stringWidth("Audi", bold, 18) + 1
        s = 12.5 / 100.0
        def poly(points, color):
            canvas.setFillColor(colors.HexColor(color))
            p = canvas.beginPath()
            pts = [(xw + px * s, y0 - 0.5 + (100 - py) * s) for px, py in points]
            p.moveTo(*pts[0])
            for pt in pts[1:]:
                p.lineTo(*pt)
            p.close()
            canvas.drawPath(p, stroke=0, fill=1)
        poly([(0, 0), (34, 0), (68, 50), (34, 100), (0, 100), (34, 50)], "#B9E10C")
        poly([(113, 0), (79, 0), (45, 50), (79, 100), (113, 100), (79, 50)], "#FFFFFF")
        canvas.setFillColor(colors.white)
        canvas.setFont(regular, 16)
        canvas.drawString(xw + 113 * s + 5, y0, "Vault")
        canvas.setFont(regular, 8.5)
        canvas.drawString(x0, h - 20 * mm, f"{FIRM} · ")
        tw = canvas.stringWidth(f"{FIRM} · ", regular, 8.5)
        canvas.drawString(x0 + tw, h - 20 * mm, "Where Audits Meet ")
        tw += canvas.stringWidth("Where Audits Meet ", regular, 8.5)
        canvas.setFillColor(colors.HexColor("#B9E10C"))
        canvas.drawString(x0 + tw, h - 20 * mm, "Exceptionalism")
        canvas.setFillColor(colors.white)
        canvas.setFont(bold, 11)
        canvas.drawRightString(w - 14 * mm, h - 11 * mm, pdf_text(f"{title}"))
        canvas.setFont(regular, 8.5)
        canvas.drawRightString(w - 14 * mm, h - 16 * mm, pdf_text(f"{client_name} · {period}"))
        canvas.drawRightString(w - 14 * mm, h - 21 * mm, generated)
        canvas.setFillColor(colors.HexColor("#5A6556"))
        canvas.setFont(regular, 7.5)
        canvas.drawString(14 * mm, 8 * mm, ASSOCIATED + f" · Powered by {FIRM}")
        canvas.drawRightString(w - 14 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=32 * mm,
                            bottomMargin=16 * mm, title=f"{title} - {client_name}", author=FIRM)
    doc.build(story_fn(), onFirstPage=letterhead, onLaterPages=letterhead)
    return buf.getvalue()


def styles():
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    regular, bold, _ = pdf_fonts()
    ss = getSampleStyleSheet()
    return {
        "h1": ParagraphStyle("h1", parent=ss["Heading1"], fontName=bold, fontSize=14, spaceAfter=6),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName=bold, fontSize=11, spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("body", parent=ss["BodyText"], fontName=regular, fontSize=8.5, leading=11),
        "bold": ParagraphStyle("bold", parent=ss["BodyText"], fontName=bold, fontSize=9, leading=12),
        "small": ParagraphStyle("small", parent=ss["BodyText"], fontName=regular, fontSize=7.5, leading=9.5,
                                textColor=colors.HexColor("#5A6556")),
        "cell": ParagraphStyle("cell", parent=ss["BodyText"], fontName=regular, fontSize=7.5, leading=9.5),
    }


def pdf_table(rows, col_widths=None, header=True, num_cols=(), footer=False):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    regular, bold, _ = pdf_fonts()
    st = styles()
    data = []
    for row in rows:
        out = []
        for c_i, v in enumerate(row):
            text = pdf_text(v)
            if isinstance(v, str) and len(text) > 30 and c_i not in num_cols:
                out.append(Paragraph(_escape(text), st["cell"]))
            else:
                out.append(text)
        data.append(out)
    t = Table(data, colWidths=col_widths, repeatRows=1 if header else 0)
    style = [
        ("FONTNAME", (0, 0), (-1, -1), regular),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DADFD2")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if header:
        style += [("FONTNAME", (0, 0), (-1, 0), bold), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F7F9F3")),
                  ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#5A6556"))]
    if footer:
        style += [("FONTNAME", (0, -1), (-1, -1), bold), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F7F9F3"))]
    for c in num_cols:
        style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(style))
    return t


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def paragraph(text, style="body"):
    from reportlab.platypus import Paragraph

    return Paragraph(_escape(pdf_text(text)), styles()[style])


def stamp() -> str:
    return datetime.now().strftime("%Y%m%d")
