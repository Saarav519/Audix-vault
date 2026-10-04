"""Stage 6: exports, emails, overview, activity log, backups."""

import gzip
import io
from datetime import date

import boto3
import pytest
from django.core.management import call_command
from django.test import Client as HttpClient
from django.test import override_settings
from django.urls import reverse
from moto import mock_aws
from openpyxl import load_workbook

from activity.models import ActivityLog
from core import backup, storage
from tests import factories as f

pytestmark = pytest.mark.django_db
LINES = [dict(stock_qty=1000, stock_value=100000, physical_qty=970, physical_value=97000, damage_qty=5,
              damage_value=500, wbc_qty=2, wbc_value=300)]


@pytest.fixture
def data(client_a, client_b, user_a, user_b):
    a1 = f.make_audit(client_a, audit_date=date(2026, 4, 10), lines=LINES)
    a2 = f.make_audit(client_a, audit_date=date(2026, 7, 20), lines=LINES, sale_value=1500000)
    b1 = f.make_audit(client_b, audit_date=date(2026, 7, 1), lines=LINES)
    return a1, a2, b1


def http_for(user):
    h = HttpClient()
    h.force_login(user)
    return h


def test_aging_excel_has_letterhead_and_numbers(data, user_a):
    r = http_for(user_a).get(reverse("portal:aging_export", args=["xlsx"]) + "?period=fy-2026")
    assert r.status_code == 200 and "attachment" in r["Content-Disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Audit register"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Audix Vault" in values and "Audix Solutions & Co" in values
    assert any(isinstance(v, str) and "Alpha Retail" in v for v in values)
    assert any(isinstance(v, str) and "Vikas Kshitij" in v for v in values)
    header_row = next(row for row in ws.iter_rows() if row[0].value == "Audit date")
    data_row = ws[header_row[0].row + 1]
    assert isinstance(data_row[3].value, (int, float))  # stock value is a real number, not text
    assert data_row[3].value == 100000.0
    assert "Quarter calendar" in wb.sheetnames
    assert ActivityLog.objects.filter(action_type="export").exists()


@pytest.mark.parametrize("url", ["aging_pdf", "audits_pdf", "audits_xlsx"])
def test_other_exports(data, user_a, url):
    names = {"aging_pdf": reverse("portal:aging_export", args=["pdf"]),
             "audits_pdf": reverse("portal:audits_export", args=["pdf"]),
             "audits_xlsx": reverse("portal:audits_export", args=["xlsx"])}
    r = http_for(user_a).get(names[url])
    assert r.status_code == 200
    if url.endswith("pdf"):
        assert r.content.startswith(b"%PDF") and len(r.content) > 1500
    else:
        ws = load_workbook(io.BytesIO(r.content)).active
        refs = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("AUD-")]
        assert len(refs) == 2  # only client A's audits


def test_compare_pdf_and_isolation(data, user_a):
    a1, a2, b1 = data
    h = http_for(user_a)
    r = h.get(reverse("portal:compare_pdf") + f"?a={a1.pk}&b={a2.pk}")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert h.get(reverse("portal:compare_pdf") + f"?a={a1.pk}&b={b1.pk}").status_code == 404


def test_client_param_ignored_for_client_users(data, user_a, client_b):
    r = http_for(user_a).get(reverse("portal:audits_export", args=["xlsx"]) + f"?client={client_b.pk}")
    ws = load_workbook(io.BytesIO(r.content)).active
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Alpha Retail" in text and "Beta Stores" not in text


def test_overview_and_activity_pages(data, admin):
    h = http_for(admin)
    r = h.get(reverse("console:overview"))
    assert r.status_code == 200
    html = r.content.decode()
    assert "Waiting for files" in html and data[0].reference in html  # no files yet
    r = h.get(reverse("console:activity") + "?type=admin")
    assert r.status_code == 200
    r = h.get(reverse("console:activity") + "?q=" + data[0].reference)
    assert r.status_code == 200


def test_monthly_summary_and_reminder(data, admin, client_a, mailoutbox):
    admin.email = "boss@audix.example"
    admin.save()
    call_command("send_monthly_summaries", "--month", "2026-07")
    to_a = [m for m in mailoutbox if client_a.email_list()[0] in m.to]
    assert len(to_a) == 1
    assert "July 2026" in to_a[0].subject
    assert "Exceptionalism" in to_a[0].body
    mailoutbox.clear()
    call_command("send_missing_file_reminders")
    assert len(mailoutbox) == 1 and mailoutbox[0].to == ["boss@audix.example"]
    assert data[0].reference in mailoutbox[0].body


def test_email_failure_never_breaks(monkeypatch, data):
    from notifications import emails

    def boom(*a, **k):
        raise OSError("smtp down")

    monkeypatch.setattr("django.core.mail.EmailMultiAlternatives.send", boom)
    assert emails.send_new_audit(data[0]) == 0


BACKUP = dict(BACKUP_S3_BUCKET_NAME="audix-backups", BACKUP_S3_REGION="us-east-1", BACKUP_S3_ENDPOINT_URL="",
              BACKUP_S3_ACCESS_KEY_ID="x", BACKUP_S3_SECRET_ACCESS_KEY="x")


def test_backup_database_and_restore(monkeypatch):
    monkeypatch.setattr(backup, "dump_database", lambda: gzip.compress(b"-- sql dump"))
    restored = {}
    monkeypatch.setattr(backup, "restore_sql", lambda data: restored.setdefault("sql", gzip.decompress(data)))
    with mock_aws(), override_settings(**BACKUP):
        s3 = boto3.client("s3", region_name="us-east-1")
        s3.create_bucket(Bucket="audix-backups")
        call_command("backup_database")
        keys = [o["Key"] for o in s3.list_objects_v2(Bucket="audix-backups")["Contents"]]
        assert len(keys) == 1 and keys[0].startswith("db/audix-") and keys[0].endswith(".sql.gz")
        from core.models import AppSettings

        assert AppSettings.objects.get(pk=1).last_backup_at is not None
        from django.core.management.base import CommandError

        with pytest.raises(CommandError):
            call_command("restore_database")  # needs --yes
        call_command("restore_database", "--yes")
        assert restored["sql"] == b"-- sql dump"


def test_sync_bucket_backup():
    with mock_aws(), override_settings(S3_BUCKET_NAME="files", S3_REGION="us-east-1", S3_ENDPOINT_URL="",
                                       S3_ACCESS_KEY_ID="x", S3_SECRET_ACCESS_KEY="x", **BACKUP):
        storage._s3.cache_clear()
        s3 = boto3.client("s3", region_name="us-east-1")
        s3.create_bucket(Bucket="files")
        s3.create_bucket(Bucket="audix-backups")
        s3.put_object(Bucket="files", Key="clients/1/a.pdf", Body=b"abc")
        s3.put_object(Bucket="files", Key="clients/1/b.jpg", Body=b"defg")
        call_command("sync_bucket_backup")
        keys = sorted(o["Key"] for o in s3.list_objects_v2(Bucket="audix-backups")["Contents"])
        assert keys == ["files/clients/1/a.pdf", "files/clients/1/b.jpg"]
        copied, skipped = backup.sync_bucket(s3, "files", s3, "audix-backups")
        assert (copied, skipped) == (0, 2)
    storage._s3.cache_clear()


def test_backup_without_bucket_fails_cleanly():
    from django.core.management.base import CommandError

    with override_settings(BACKUP_S3_BUCKET_NAME=""):
        with pytest.raises(CommandError):
            call_command("backup_database")
