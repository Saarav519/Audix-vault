"""Stage 7: old-data import round trip and seed_demo."""

import io
from datetime import date

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.urls import reverse
from openpyxl import Workbook

from audits import importer
from audits.models import Audit
from clients.models import Client

pytestmark = pytest.mark.django_db


def xlsx(rows):
    wb = Workbook()
    wb.active.append(importer.COLUMNS)
    for r in rows:
        wb.active.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def row(code="S01", d="2025-04-05", cat="Grocery & staples", t="Physical stock audit", sale=500000, stock=100000):
    return ["Alpha Retail", code, d, t, "Morning", sale, cat, 1000, stock, 960, 96000, 10, 1000, 5, 500]


def test_template_downloads(admin_client):
    r = admin_client.get(reverse("console:import_template"))
    assert r.status_code == 200
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == importer.COLUMNS


def test_import_round_trip(admin_client, client_a):
    data = xlsx([row(), row(cat="Dairy & frozen", stock=50000), row(d="2025-07-10"), row(d="2025-07-10", cat="Beverages")])
    r = admin_client.post(reverse("console:import"), {"file": SimpleUploadedFile("old.xlsx", data), "dry_run": "1"})
    assert r.status_code == 200 and r.context["result"].ok
    assert not Audit.objects.exists()  # dry run saves nothing
    admin_client.post(reverse("console:import"), {"file": SimpleUploadedFile("old.xlsx", data)})
    audits = list(Audit.objects.order_by("audit_date"))
    assert len(audits) == 2
    first, second = audits
    assert first.is_imported and first.is_published and first.lines.count() == 2
    assert first.observations.get().text == "No observation recorded"
    assert second.previous_audit_id == first.pk and second.aging_days == 96 and second.delay_days == 6
    assert first.stock_value == 150000
    # importing again reports duplicates
    res = importer.parse("old.xlsx", data)
    assert res.errors and "Duplicate" in res.errors[0][1]


def test_import_reports_row_errors(client_a):
    data = xlsx([row(code="ZZ9"), row(cat="Unknown cat"), ["Alpha Retail", "S01", "not a date", "Physical stock audit",
                                                           "Morning", "", "Beverages", 1, 1, 1, 1, 0, 0, 0, 0],
                 ["Alpha Retail", "S01", "2025-04-05", "Physical stock audit", "Morning", "", "Beverages", -1, 1, 1, 1, 0, 0, 0, 0],
                 ["Nobody Ltd"] + row()[1:]])
    res = importer.parse("x.xlsx", data)
    msgs = " | ".join(m for _, m in res.errors)
    assert "Unknown store code" in msgs and "Unknown category" in msgs and "Bad audit date" in msgs
    assert "must be a number" in msgs and "Unknown client" in msgs
    assert [n for n, _ in res.errors] == [2, 3, 4, 5, 6]


def test_csv_import(client_a, admin):
    text = ",".join(importer.COLUMNS) + "\n" + ",".join(str(x) for x in row()) + "\n"
    res = importer.parse("old.csv", text.encode())
    assert res.ok
    assert importer.run_import(res, admin) == 1


def test_seed_demo(monkeypatch, admin):
    monkeypatch.setenv("DEMO_CLIENT_PASSWORD", "Demo-pass-12345")
    from django.utils import timezone

    monkeypatch.setattr(timezone, "localdate", lambda: date(2024, 12, 31))
    call_command("seed_demo", "--no-files")
    green = Client.objects.get(slug="greenfield-retail")
    assert green.stores.count() == 8 and green.categories.count() == 7
    assert Client.objects.filter(slug="urban-mart").exists()
    audits = Audit.objects.filter(client=green)
    assert audits.count() > 50
    assert not audits.filter(store__code="GR07", sale_value__isnull=False).exists()
    assert audits.filter(series="full", delay_days__gt=0).exists()
    assert all(a.observations.exists() for a in audits[:20])
    call_command("seed_demo", "--no-files")  # idempotent
    assert Client.objects.filter(slug="greenfield-retail").count() == 1
    from accounts.models import User

    assert User.objects.get(login_id="greenfield").check_password("Demo-pass-12345")


def test_seed_demo_refuses_when_not_allowed(settings, monkeypatch):
    settings.DEBUG = False
    settings.ALLOW_DEMO_SEED = False
    call_command("seed_demo", "--if-allowed")
    assert not Client.objects.exists()


