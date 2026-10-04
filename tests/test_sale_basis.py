"""'% of sale value' must reconcile with the rupees and sale total shown next to it."""

import io
from datetime import date
from decimal import Decimal

import pytest
from django.test import Client as HttpClient
from django.urls import reverse
from openpyxl import load_workbook

from core import calc
from core.formatting import inr, pct
from tests import factories as f

TODAY = date(2026, 10, 4)


def L(stock, diff):
    return [dict(stock_qty=stock / 100, stock_value=stock, physical_qty=(stock + diff) / 100,
                 physical_value=stock + diff)]


# ---------------------------------------------------------------- calc


def test_sale_basis_partial():
    audits = [
        {"diff_value": -5000, "stock_value": 100000, "sale_value": 1000000},
        {"diff_value": -3000, "stock_value": 100000, "sale_value": 500000},
        {"diff_value": -2000, "stock_value": 100000, "sale_value": None},
        {"diff_value": -1000, "stock_value": 100000, "sale_value": 0},
    ]
    b = calc.sale_basis(audits)
    assert (b.n_total, b.n_with_sale) == (4, 2)
    assert b.net_diff_all == -11000 and b.stock_total == 400000
    assert b.pct_stock_all == Decimal("-2.75")
    assert b.net_diff_with_sale == -8000 and b.sale_total == 1500000
    assert b.pct_sale == Decimal(-8000) / Decimal(1500000) * 100
    assert b.partial_sale and not b.all_have_sale
    assert b.sale_text() == "−0.53% of sale value (−₹8,000 on the 2 of 4 audits that have a sale value)"
    assert b.stock_text() == "−2.75% of stock value"


def test_sale_basis_all_and_none():
    all_sale = calc.sale_basis([{"diff_value": -100, "stock_value": 1000, "sale_value": 10000}] * 2)
    assert all_sale.all_have_sale and all_sale.sale_text() == "−1.00% of sale value"
    none = calc.sale_basis([{"diff_value": -100, "stock_value": 1000, "sale_value": None}])
    assert not none.has_any_sale and none.sale_text() is None and none.pct_sale is None
    empty = calc.sale_basis([])
    assert empty.n_total == 0 and empty.pct_stock_all == 0 and empty.sale_text() is None


def test_from_totals_matches_rows():
    rows = [{"diff_value": -700, "stock_value": 9000, "sale_value": 50000},
            {"diff_value": 200, "stock_value": 4000, "sale_value": None}]
    a = calc.sale_basis(rows)
    b = calc.sale_basis_from_totals(2, 1, -500, 13000, -700, 50000)
    assert (a.pct_sale, a.pct_stock_all, a.sale_text()) == (b.pct_sale, b.pct_stock_all, b.sale_text())


# ---------------------------------------------------------------- screens and exports


@pytest.fixture
def mixed(db, client_a, user_a, monkeypatch):
    """Three full audits in FY 2026-27: two with a sale value, one without."""
    monkeypatch.setattr("reports.views.today", lambda: TODAY)
    s1, s2 = client_a.stores.all()
    f.make_audit(client_a, store=s1, audit_date=date(2026, 5, 1), lines=L(200000, -6000), sale_value=1200000)
    f.make_audit(client_a, store=s1, audit_date=date(2026, 8, 1), lines=L(100000, -2000), sale_value=800000)
    f.make_audit(client_a, store=s2, audit_date=date(2026, 9, 1), lines=L(300000, -9000))  # no sale value
    h = HttpClient()
    h.force_login(user_a)
    return h


@pytest.mark.django_db
def test_aging_kpi_and_footer_reconcile(mixed):
    r = mixed.get(reverse("portal:aging") + "?period=fy-2026")
    k, foot = r.context["report"].kpis, r.context["report"].footer
    assert k["net_diff"] == -17000
    # sale % uses only the audits that have a sale value, and says so
    assert foot["n_sale"] == 2 and foot["n"] == 3
    assert foot["diff_sale"] == -8000 and foot["sale"] == 2000000
    assert foot["pct_sale"] == Decimal("-0.4")  # -8,000 / 20,00,000
    assert foot["diff_sale"] / foot["sale"] * 100 == foot["pct_sale"]
    assert foot["diff"] / foot["stock"] * 100 == foot["pct_stock"]  # all audits on the stock base
    assert k["net_text"] == "−0.40% of sale value (−₹8,000 on the 2 of 3 audits that have a sale value)"
    html = r.content.decode()
    assert "All audits (3)" in html and "Audits with sale value (2 of 3)" in html
    assert "₹20.00 L" in html  # sale total of the 2 audits
    assert k["net_text"] in html


@pytest.mark.django_db
def test_aging_excel_and_pdf_totals(mixed):
    r = mixed.get(reverse("portal:aging_export", args=["xlsx"]) + "?period=fy-2026")
    ws = load_workbook(io.BytesIO(r.content))["Audit register"]
    rows = {row[0].value: [c.value for c in row] for row in ws.iter_rows() if isinstance(row[0].value, str)}
    all_row, sale_row = rows["All audits (3)"], rows["Audits with sale value (2 of 3)"]
    assert all_row[5] == -17000 and all_row[7] == "of stock value"
    assert all_row[6] == pytest.approx(-17000 / 600000)
    assert sale_row[4] == 2000000 and sale_row[5] == -8000 and sale_row[7] == "of sale value"
    assert sale_row[6] == pytest.approx(sale_row[5] / sale_row[4])  # stored as a fraction for Excel %
    kpi_sale = rows["Net difference, audits with sale value (2 of 3)"]
    assert kpi_sale[1] == -8000 and kpi_sale[2] == pytest.approx(-0.004)
    pdf = mixed.get(reverse("portal:aging_export", args=["pdf"]) + "?period=fy-2026")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


@pytest.mark.django_db
def test_dashboard_variance_kpi_reconciles(mixed):
    r = mixed.get(reverse("portal:dashboard") + "?period=yearly")
    cur = r.context["cur"]
    b = cur["basis"]
    assert (b.n_total, b.n_with_sale) == (3, 2)
    assert b.net_diff_with_sale / b.sale_total * 100 == cur["var_pct_sale"]
    variance = next(k for k in r.context["kpis"] if k["label"] == "Variance")
    assert variance["sub"] == f"of stock value; {pct(b.pct_sale)} of sale value ({inr(-8000)} on the 2 of 3 audits that have a sale value)"
    html = r.content.decode()
    assert "on the 2 of 3 audits that have a sale value" in html  # hero line too


@pytest.mark.django_db
def test_all_audits_with_sale_keep_short_text(client_a, user_a, monkeypatch):
    monkeypatch.setattr("reports.views.today", lambda: TODAY)
    f.make_audit(client_a, audit_date=date(2026, 9, 1), lines=L(100000, -1000), sale_value=500000)
    h = HttpClient()
    h.force_login(user_a)
    k = h.get(reverse("portal:aging") + "?period=fy-2026").context["report"].kpis
    assert k["net_text"] == "−0.20% of sale value"
