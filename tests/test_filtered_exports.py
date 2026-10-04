"""Filtered audits exports (portal and admin) and the admin Audits filters."""

import io
from datetime import date

import pytest
from django.test import Client as HttpClient
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader

from activity.models import ActivityLog
from tests import factories as f

pytestmark = pytest.mark.django_db
TODAY = date(2026, 10, 4)


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setattr("reports.views.today", lambda: TODAY)


@pytest.fixture
def data(client_a, client_b):
    s1, s2 = client_a.stores.order_by("code")
    out = {"s1": s1, "s2": s2}
    out["s1_audits"] = [f.make_audit(client_a, store=s1, audit_date=date(2026, m, 5)) for m in (4, 6, 8)]
    out["s1_draft"] = f.make_audit(client_a, store=s1, audit_date=date(2026, 9, 5), publish=False)
    out["s2_audits"] = [f.make_audit(client_a, store=s2, audit_date=date(2026, m, 9), audit_type="cycle")
                        for m in (5, 7)]
    out["b"] = f.make_audit(client_b, audit_date=date(2026, 9, 20))
    return out


def http_for(user):
    h = HttpClient()
    h.force_login(user)
    return h


def xlsx_rows(resp):
    ws = load_workbook(io.BytesIO(resp.content)).active
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def xlsx_refs(rows, ref_col=1):
    return {r[ref_col] for r in rows if isinstance(r[ref_col], str) and r[ref_col].startswith("AUD-")}


def pdf_text(resp):
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(resp.content)).pages)


def flat(text):
    return " ".join(text.split())


# ---------------------------------------------------------------- portal


def test_portal_store_filter_export_matches_screen(data, user_a):
    h = http_for(user_a)
    qs = f"?store={data['s1'].pk}"
    screen = h.get(reverse("portal:audits") + qs)
    shown = {a.reference for a in screen.context["page"]}
    count = screen.context["page"].paginator.count
    assert count == 3  # drafts never reach the portal

    rows = xlsx_rows(h.get(reverse("portal:audits_export", args=["xlsx"]) + qs))
    assert xlsx_refs(rows) == shown
    header = "\n".join(str(r[0]) for r in rows[:6])
    assert "Store: S01" in header and data["s1"].name in header
    filters = next(r for r in rows if r[0] == "Filters")
    assert f"Store: S01 · {data['s1'].name}" in filters[1] and "Period: Any time" in filters[1]
    assert next(r for r in rows if r[0] == "Audits")[1] == f"{count} audits"

    text = flat(pdf_text(h.get(reverse("portal:audits_export", args=["pdf"]) + qs)))
    assert f"Store: S01 · {data['s1'].name}" in text
    assert f"{count} audits" in text
    for ref in shown:
        assert ref in text
    for a in data["s2_audits"]:
        assert a.reference not in text


def test_portal_export_describes_every_filter(data, user_a):
    h = http_for(user_a)
    qs = f"?store={data['s2'].pk}&type=cycle&shift=morning&period=fy&q=S02"
    rows = xlsx_rows(h.get(reverse("portal:audits_export", args=["xlsx"]) + qs))
    line = next(r for r in rows if r[0] == "Filters")[1]
    for part in ("Store: S02", "Audit type: Cycle count", "Shift: Morning", "Period: This financial year",
                 "Search: “S02”"):
        assert part in line, part


def test_portal_export_crosses_pages(client_a, user_a):
    for i in range(30):
        f.make_audit(client_a, audit_date=date(2026, 9, 1 + i % 28))
    rows = xlsx_rows(http_for(user_a).get(reverse("portal:audits_export", args=["xlsx"])))
    assert len(xlsx_refs(rows)) == 30


def test_export_cap_is_stated(data, user_a, monkeypatch):
    monkeypatch.setattr("reports.exports.EXPORT_CAP", 2)
    h = http_for(user_a)
    rows = xlsx_rows(h.get(reverse("portal:audits_export", args=["xlsx"])))
    assert len(xlsx_refs(rows)) == 2
    assert "Showing the first 2 of 5 audits" in next(r for r in rows if r[0] == "Audits")[1]
    assert "Showing the first 2 of 5 audits" in flat(pdf_text(h.get(reverse("portal:audits_export", args=["pdf"]))))


# ---------------------------------------------------------------- admin list


@pytest.mark.parametrize("who", ["admin", "auditor"])
def test_admin_store_filter_and_export(data, client_a, request, who):
    h = http_for(request.getfixturevalue(who))
    qs = f"?client={client_a.pk}&store={data['s1'].pk}"
    screen = h.get(reverse("audits:list") + qs)
    assert screen.status_code == 200
    shown = {a.reference for a in screen.context["page"]}
    count = screen.context["page"].paginator.count
    assert count == 4 and data["s1_draft"].reference in shown  # admin sees drafts too
    html = screen.content.decode()
    assert reverse("audits:export", args=["xlsx"]) in html and reverse("audits:export", args=["pdf"]) in html

    rows = xlsx_rows(h.get(reverse("audits:export", args=["xlsx"]) + qs))
    assert xlsx_refs(rows) == shown
    line = next(r for r in rows if r[0] == "Filters")[1]
    assert "Client: Alpha Retail" in line and f"Store: S01 · {data['s1'].name}" in line
    assert next(r for r in rows if r[0] == "Audits")[1] == f"{count} audits"
    assert "Draft" in {r[-1] for r in rows}

    text = flat(pdf_text(h.get(reverse("audits:export", args=["pdf"]) + qs)))
    assert f"Store: S01 · {data['s1'].name}" in text and f"{count} audits" in text
    assert all(ref in text for ref in shown)
    assert not any(a.reference in text for a in data["s2_audits"])
    assert ActivityLog.objects.filter(action_type="export", client=client_a).count() == 2


def test_admin_filters_type_shift_period_status(data, admin_client, client_a):
    url = reverse("audits:list")
    r = admin_client.get(url + "?type=cycle")
    assert {a.reference for a in r.context["page"]} == {a.reference for a in data["s2_audits"]}
    r = admin_client.get(url + "?status=draft")
    assert [a.reference for a in r.context["page"]] == [data["s1_draft"].reference]
    r = admin_client.get(url + "?period=30")
    assert {a.reference for a in r.context["page"]} == {data["s1_draft"].reference, data["b"].reference}
    r = admin_client.get(url + "?shift=night")
    assert r.context["page"].paginator.count == 0
    r = admin_client.get(url + "?q=Beta")  # search still matches the client name
    assert [a.reference for a in r.context["page"]] == [data["b"].reference]


def test_admin_store_needs_client(data, admin_client, client_b):
    url = reverse("audits:list")
    r = admin_client.get(url)
    assert "Choose a client first" in r.content.decode()
    # a store from another client is ignored rather than emptying the list
    r = admin_client.get(url + f"?client={client_b.pk}&store={data['s1'].pk}")
    assert r.context["f"]["store"] == "" and r.context["page"].paginator.count == 1


def test_admin_export_all_clients_has_client_column(data, admin_client):
    rows = xlsx_rows(admin_client.get(reverse("audits:export", args=["xlsx"])))
    head = next(r for r in rows if r[0] == "Date")
    assert head[2] == "Client"
    assert "Client: All clients" in next(r for r in rows if r[0] == "Filters")[1]
    assert len(xlsx_refs(rows)) == 7


def test_admin_export_permissions(data, user_a, client):
    assert http_for(user_a).get(reverse("audits:export", args=["xlsx"])).status_code in (302, 403, 404)
    r = client.get(reverse("audits:export", args=["xlsx"]))
    assert r.status_code == 302 and "login" in r["Location"]
    assert http_for(user_a).get(reverse("audits:export", args=["csv"])).status_code in (302, 403, 404)
