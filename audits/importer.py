"""Import historical audits from Excel or CSV (one row per audit per category), with a dry run."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

from audits import services
from audits.models import Audit, AuditLine, AuditStatus, AuditType, Observation, Shift
from clients.models import Category, Client, Store
from core import calc

COLUMNS = ["client", "store_code", "audit_date", "audit_type", "shift", "sale_value", "category", "stock_qty",
           "stock_value", "physical_qty", "physical_value", "damage_qty", "damage_value", "wbc_qty", "wbc_value"]
TYPE_LOOKUP = {**{v: v for v in AuditType.values}, **{label.lower(): v for v, label in AuditType.choices},
               "full": "physical", "physical": "physical"}
SHIFT_LOOKUP = {**{v: v for v in Shift.values}, **{label.lower(): v for v, label in Shift.choices}}
MAX_ROWS = 50000


@dataclass
class ImportResult:
    errors: list = field(default_factory=list)  # (row number, message)
    audits: list = field(default_factory=list)  # planned audits (dicts)
    rows: int = 0
    created: int = 0

    @property
    def ok(self):
        return not self.errors and bool(self.audits)


def template_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Audits"
    ws.append(COLUMNS)
    ws.append(["Greenfield Retail", "GR01", "2025-04-05", "Physical stock audit", "Morning", 2500000,
               "Grocery & staples", 1000, 100000, 960, 96000, 10, 1000, 5, 500])
    ws.append(["Greenfield Retail", "GR01", "2025-04-05", "Physical stock audit", "Morning", 2500000,
               "Beverages", 400, 40000, 398, 39800, 0, 0, 0, 0])
    notes = wb.create_sheet("How to fill")
    for line in [
        "One row per audit per category. Rows with the same client, store_code, audit_date, audit_type and shift form one audit.",
        "client: the client name exactly as in Audix Vault.",
        "audit_date: YYYY-MM-DD (or an Excel date).",
        "audit_type: Physical stock audit, Surprise audit, Cycle count or Closing audit.",
        "shift: Morning, Afternoon or Night.",
        "sale_value: optional; sales since the previous audit of the same store. Use the same value on every row of the audit.",
        "Numbers must be zero or more. Leave a cell empty for zero.",
    ]:
        notes.append([line])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def read_rows(name: str, data: bytes):
    """Yield (row_number, dict) from an .xlsx or .csv file."""
    if name.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        header = None
        for i, row in enumerate(reader, start=1):
            if header is None:
                header = [h.strip().lower() for h in row]
                continue
            if any(c.strip() for c in row):
                yield i, dict(zip(header, row, strict=False))
        return
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    header = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if header is None:
            header = [str(h or "").strip().lower() for h in row]
            continue
        if any(c not in (None, "") for c in row):
            yield i, dict(zip(header, row, strict=False))


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _num(v, places):
    if v is None or str(v).strip() == "":
        return Decimal(0)
    try:
        d = Decimal(str(v).replace(",", "").replace("₹", "").strip())
    except InvalidOperation:
        return None
    if not d.is_finite() or d < 0:
        return None
    return d.quantize(Decimal(1).scaleb(-places))


def parse(name: str, data: bytes) -> ImportResult:
    res = ImportResult()
    clients = {c.name.lower(): c for c in Client.objects.all()}
    clients.update({c.slug: c for c in clients.values()})
    stores, cats = {}, {}
    groups = {}
    for n, row in read_rows(name, data):
        res.rows += 1
        if res.rows > MAX_ROWS:
            res.errors.append((n, f"Too many rows (limit {MAX_ROWS})."))
            break
        missing = [c for c in COLUMNS if c not in row]
        if missing:
            res.errors.append((n, f"Missing columns: {', '.join(missing)}. Use the template."))
            break
        client = clients.get(str(row["client"] or "").strip().lower())
        if client is None:
            res.errors.append((n, f"Unknown client “{row['client']}”."))
            continue
        if client.pk not in stores:
            stores[client.pk] = {s.code.upper(): s for s in Store.objects.filter(client=client)}
            cats[client.pk] = {c.name.lower(): c for c in Category.objects.filter(client=client)}
        store = stores[client.pk].get(str(row["store_code"] or "").strip().upper())
        if store is None:
            res.errors.append((n, f"Unknown store code “{row['store_code']}” for {client.name}."))
            continue
        cat = cats[client.pk].get(str(row["category"] or "").strip().lower())
        if cat is None:
            res.errors.append((n, f"Unknown category “{row['category']}” for {client.name}."))
            continue
        d = _date(row["audit_date"])
        if d is None or d > timezone.localdate():
            res.errors.append((n, f"Bad audit date “{row['audit_date']}”."))
            continue
        atype = TYPE_LOOKUP.get(str(row["audit_type"] or "").strip().lower())
        if atype is None:
            res.errors.append((n, f"Unknown audit type “{row['audit_type']}”."))
            continue
        shift = SHIFT_LOOKUP.get(str(row["shift"] or "morning").strip().lower())
        if shift is None:
            res.errors.append((n, f"Unknown shift “{row['shift']}”."))
            continue
        nums = {}
        bad = False
        for f in calc.FIELDS:
            v = _num(row.get(f), 3 if f.endswith("qty") else 2)
            if v is None:
                res.errors.append((n, f"{f} must be a number, zero or more."))
                bad = True
                break
            nums[f] = v
        sale = _num(row.get("sale_value"), 2)
        if sale is None:
            res.errors.append((n, "sale_value must be a number, zero or more."))
            bad = True
        if bad:
            continue
        key = (client.pk, store.pk, d, atype, shift)
        g = groups.setdefault(key, {"client": client, "store": store, "date": d, "type": atype, "shift": shift,
                                    "sale": sale or None, "lines": {}, "row": n})
        if (sale or None) != g["sale"]:
            res.errors.append((n, "sale_value differs between rows of the same audit."))
            continue
        if cat.pk in g["lines"]:
            res.errors.append((n, f"Category {cat.name} appears twice for this audit."))
            continue
        g["lines"][cat.pk] = (cat, nums)
    for g in groups.values():
        if Audit.objects.live().filter(client=g["client"], store=g["store"], audit_date=g["date"],
                                       audit_type=g["type"]).exists():
            res.errors.append((g["row"], f"Duplicate: {g['store'].code} already has a "
                                         f"{calc.AUDIT_TYPE_LABELS[g['type']].lower()} on {g['date']:%d %b %Y}."))
        else:
            res.audits.append(g)
    return res


@transaction.atomic
def run_import(res: ImportResult, user) -> int:
    chains = set()
    for g in sorted(res.audits, key=lambda x: x["date"]):
        audit = Audit.objects.create(client=g["client"], store=g["store"], audit_date=g["date"], audit_type=g["type"],
                                     shift=g["shift"], sale_value=g["sale"], is_imported=True,
                                     status=AuditStatus.PUBLISHED, created_by=user, published_by=user,
                                     published_at=timezone.now())
        for cat, nums in g["lines"].values():
            AuditLine.objects.create(audit=audit, category=cat, **nums)
        Observation.objects.create(audit=audit, category=None, kind="process_gap", severity="low",
                                   text="No observation recorded", recommendation="")
        audit.refresh_snapshot()
        chains.add((g["client"], audit.store_id, audit.series))
        res.created += 1
    for client, store_id, series in chains:
        services.relink_series(client, store_id, series)
    return res.created
