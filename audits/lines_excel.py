"""Category-wise summary as an Excel file: a blank template to fill in, and reading a filled one back.

The upload only fills the Add/Edit grid in the browser; nothing is saved until the form is saved,
and every number still goes through the same checks as typed numbers.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from audits.entry import FIELD_LABELS, parse_decimal
from core import calc
from core import exports as ex

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
HEADERS = ["Category"] + [FIELD_LABELS[f] for f in calc.FIELDS]
TOTAL_WORDS = {"total", "storetotal", "grandtotal"}


def _norm(text) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").casefold())


HEADER_KEYS = {_norm(FIELD_LABELS[f]): f for f in calc.FIELDS}
HEADER_KEYS.update({_norm(FIELD_LABELS[f].replace("qty", "quantity")): f for f in calc.FIELDS})


def template_xlsx(client, categories) -> bytes:
    """Blank template in the grid's layout, one row per category of the client."""
    from openpyxl.styles import Alignment, Font, PatternFill

    wb, ws = ex.new_workbook()
    ws.title = "Category summary"
    ws.append(HEADERS)
    for i, _ in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F7F9F3")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for cat in categories:
        ws.append([cat.name] + [None] * len(calc.FIELDS))
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 14
    ws.freeze_panes = "B2"
    help_ws = wb.create_sheet("How to fill")
    for line in (
        f"Audix Vault · category-wise summary · {client.name}",
        "",
        "Fill one row per category on the first sheet. Keep the header row and the category names as they are.",
        "Quantities may have up to 3 decimals; values are in rupees, up to 2 decimals. Leave a cell blank for zero.",
        "Total physical, difference and variance are worked out by Audix Vault, so they are not in the file.",
        "Then use 'Upload Excel' in Step 2 of Add audit. The numbers fill the grid; check them and save.",
    ):
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 110
    return ex.workbook_bytes(wb)


@dataclass
class ParseResult:
    values: dict = field(default_factory=dict)  # category id -> {field: text}
    matched: list = field(default_factory=list)  # category names
    unknown: list = field(default_factory=list)  # names in the file that are not categories of the client
    missing: list = field(default_factory=list)  # categories of the client not in the file
    errors: list = field(default_factory=list)

    def as_json(self):
        return {"values": self.values, "matched": self.matched, "unknown": self.unknown, "missing": self.missing,
                "errors": self.errors}


def _rows(upload):
    name = (upload.name or "").lower()
    if name.endswith(".csv"):
        import csv
        import io

        text = upload.read().decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(text)))
    if not name.endswith((".xlsx", ".xlsm")):
        raise ValueError("Upload an Excel file (.xlsx) or a CSV file. Old .xls files: save them as .xlsx first.")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        raise ValueError("This file could not be read as an Excel workbook.")
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS)]


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse(upload, categories) -> ParseResult:
    """Read the first sheet. The header row is the first row with a 'Category' cell in the first
    ten rows; columns are matched by their header text, so their order does not matter."""
    res = ParseResult()
    if upload.size > MAX_BYTES:
        res.errors.append("The file is larger than 5 MB.")
        return res
    try:
        rows = _rows(upload)
    except ValueError as e:
        res.errors.append(str(e))
        return res
    header_idx, cols = None, {}
    for i, row in enumerate(rows[:10]):
        keys = [_norm(c) for c in row]
        if "category" in keys:
            header_idx = i
            cols = {"category": keys.index("category")}
            for j, k in enumerate(keys):
                if k in HEADER_KEYS:
                    cols[HEADER_KEYS[k]] = j
            break
    if header_idx is None:
        res.errors.append("No header row with a 'Category' column was found. Use the template from 'Download template'.")
        return res
    absent = [FIELD_LABELS[f] for f in calc.FIELDS if f not in cols]
    if absent:
        res.errors.append("These columns are missing: " + ", ".join(absent) + ".")
        return res

    by_name = {_norm(c.name): c for c in categories}
    seen = set()
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        cells = list(row) + [None] * (max(cols.values()) + 1 - len(row))
        name = _cell_text(cells[cols["category"]])
        nums = {f: _cell_text(cells[cols[f]]) for f in calc.FIELDS}
        if not name:
            if any(nums.values()):
                res.errors.append(f"Row {n}: numbers without a category name.")
            continue
        if _norm(name) in TOTAL_WORDS:
            continue
        cat = by_name.get(_norm(name))
        if cat is None:
            res.unknown.append(name)
            continue
        if cat.pk in seen:
            res.errors.append(f"Row {n}: {cat.name} appears more than once.")
            continue
        seen.add(cat.pk)
        out = {}
        for f in calc.FIELDS:
            v, err = parse_decimal(nums[f].replace("−", "-"), 3 if f.endswith("qty") else 2)
            if err:
                res.errors.append(f"Row {n}, {cat.name}: {FIELD_LABELS[f]} {err}.")
                out[f] = nums[f]
            else:
                out[f] = "" if nums[f] == "" else format(v.normalize(), "f")
        res.values[str(cat.pk)] = out
        res.matched.append(cat.name)
    res.missing = [c.name for c in categories if c.pk not in seen]
    if not res.matched and not res.errors:
        res.errors.append("No row in the file matches a category of this client.")
    return res
